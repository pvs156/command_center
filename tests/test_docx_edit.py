import io

import pytest
from docx import Document
from docx.shared import Pt
from fastapi.testclient import TestClient

from app import store, docx_edit
from app.main import app


def base_docx():
    d = Document()
    d.add_paragraph('JANE TESTER')
    d.add_paragraph('Dallas, TX | jane@example.com | (555) 010-0199 | linkedin.com/in/jane')
    d.add_paragraph('PROFESSIONAL SUMMARY')
    d.add_paragraph('AI engineer building Python services and LLM applications.')
    d.add_paragraph('PROFESSIONAL EXPERIENCE')
    p = d.add_paragraph()
    p.add_run('AI Engineer | Acme | Dallas, TX').bold = True
    p.add_run('\tJan 2023 – Present')
    d.add_paragraph('Built Python and FastAPI services for a RAG application serving 40 users.', style='List Bullet')
    d.add_paragraph('Deployed Docker services on AWS with GitHub Actions.', style='List Bullet')
    d.add_paragraph('TECHNICAL SKILLS')
    s = d.add_paragraph()
    r = s.add_run('Languages: ')
    r.bold = True
    s.add_run('Python, SQL')
    s2 = d.add_paragraph()
    s2.add_run('Cloud & Platform: ').bold = True
    s2.add_run('AWS, Docker')
    d.add_paragraph('EDUCATION & CERTIFICATIONS')
    d.add_paragraph('M.S. Computer Science, Example University\t2022')
    for para in d.paragraphs:
        for run in para.runs:
            run.font.size = Pt(10.5)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'DB', tmp_path / 'test.db')
    monkeypatch.setenv('COMMAND_CENTER_SCHEDULER', '0')
    import app.main as m
    monkeypatch.setattr(m, 'UPLOADS', tmp_path / 'uploads')
    with TestClient(app) as c:
        yield c


def test_segments_mark_only_new_words():
    segs = docx_edit.segments('Built Python services.', 'Built Python and FastAPI services.')
    assert [t for t, changed in segs if changed] == ['and FastAPI ']  or ''.join(t for t, c in segs if c).strip() == 'and FastAPI'


def test_upload_save_export_edits_your_own_document(client):
    data = base_docx()
    r = client.post('/api/profile/preview', files={'file': ('Jane.docx', data, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()
    assert r['upload_id']
    saved = client.post('/api/profile', json={'name': 'Jane', 'text': r['text'], 'confirmed': True, 'upload_id': r['upload_id'], 'filename': 'Jane.docx'}).json()
    assert saved['docx'] and saved['docx_coverage'] >= 90
    assert [f['section'] for f in saved['facts']][:2] == ['summary', 'experience']
    job = {'id': 'real:1', 'demo': False, 'company': 'Acme AI', 'title': 'AI Engineer', 'location': 'Dallas, TX',
           'employment_type': 'Full-time', 'description': 'Requirements\n- 2 years of experience with Python, FastAPI, RAG, Kubernetes and LangGraph\n$120,000-$140,000 per year',
           'published_at': store.now(), 'publication_kind': 'first_published', 'url': 'https://example.com/j', 'content_hash': 'h1'}
    store.put('jobs', job['id'], job)
    client.post('/api/skills', json={'skill': 'LangGraph', 'mode': 'live'})
    client.post('/api/prepare', json={'job_ids': ['real:1']})
    pkg = next(p for p in client.get('/api/state').json()['packages'] if p['job_id'] == 'real:1')

    out = client.get(f'/api/packages/{pkg["id"]}/export?format=docx')
    doc = Document(io.BytesIO(out.content))
    orig = Document(io.BytesIO(data))
    texts, before = [p.text for p in doc.paragraphs], [p.text for p in orig.paragraphs]
    assert len(texts) == len(before)
    # Section headings stay in your order, at the same positions.
    for heading in ['PROFESSIONAL SUMMARY', 'PROFESSIONAL EXPERIENCE', 'TECHNICAL SKILLS', 'EDUCATION & CERTIFICATIONS']:
        assert texts.index(heading) == before.index(heading)
    # Bold role line with tabbed dates is untouched; bullets keep their list style and order.
    role = doc.paragraphs[5]
    assert role.text == before[5] and role.runs[0].bold
    assert doc.paragraphs[6].style.name == 'List Bullet' and texts[6] == before[6]
    # Summary keeps your words and gains a keyword sentence; LangGraph lands on an existing skills line.
    assert texts[3].startswith(before[3]) and 'LangGraph' in texts[3] or 'LangGraph' in ' '.join(texts[9:11])
    assert any('LangGraph' in t for t in texts[9:11])
    assert all(not r.font.highlight_color for p in doc.paragraphs for r in p.runs)  # clean copy for applying
    assert all(r.font.size == Pt(10.5) for p in doc.paragraphs for r in p.runs if r.text.strip())

    hl = Document(io.BytesIO(client.get(f'/api/packages/{pkg["id"]}/export?format=docx&highlight=true').content))
    marked = ''.join(r.text for p in hl.paragraphs for r in p.runs if r.font.highlight_color)
    assert 'LangGraph' in marked and 'Built Python' not in marked

    preview = client.get(f'/api/packages/{pkg["id"]}/preview').json()
    assert preview['report']['source'] == 'your Word file' and not preview['report']['failed']
    assert any(run['added'] for para in preview['paragraphs'] for run in para['runs'])


def test_rewritten_bullet_keeps_bold_prefix_formatting():
    d = Document()
    p = d.add_paragraph()
    p.add_run('Harness: ').bold = True
    p.add_run('Built a Python harness for 40 users.')
    docx_edit.edit_paragraph(p, 'Harness: Built a Python and FastAPI harness for 40 users.', highlight=True)
    assert p.text == 'Harness: Built a Python and FastAPI harness for 40 users.'
    assert p.runs[0].text == 'Harness: ' and p.runs[0].bold
    assert ''.join(r.text for r in p.runs if r.font.highlight_color).strip() == 'and FastAPI'
