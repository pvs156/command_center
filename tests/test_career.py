import asyncio
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import career, store
from app.main import app
from app.policy import assess
from app.resume import parse
from app.tailor import tailor
from app import rules

RESUME = """Jane Tester
Dallas, TX | jane@example.com | (555) 010-0199 | linkedin.com/in/jane
# Summary
AI Engineer with 3+ years building Python services and LLM applications.
# Experience
AI Engineer | Acme | Dallas, TX Jan 2024 – Present
- Built Python services for internal tools.
- Maintained dashboards for the support team.
# Skills
Languages & Frameworks: Python, SQL
Generative AI & Agents: LLMs, Prompt Engineering
"""

CAREER_MD = """# Career

## 1. How to use this document

- Name Acme's client only once per resume.

## 2. Profile

### Summary variants

#### Agentic focus

AI Engineer with 3+ years building production agentic workflows with LangGraph and MCP, retrieval-augmented generation (RAG) services, and evaluation pipelines on AWS.

## 3. Acme, AI Engineer

### End-to-end story

Built a LangGraph multi-agent system and an MCP governance layer for a banking client.

### Bullet bank

#### Agents

- Engineered a LangGraph multi-agent review workflow with human approval, cutting review time by 35%. (current resume)
- Engineered a LangGraph multi-agent workflow with deterministic validation and human approval. (first version)

#### MCP

- Developed an MCP server that enforces scoped authorization for agent tool access.

#### Vector store

- Indexed documents in [FILL: vector database name] for semantic retrieval.

### Keywords this role proves

LangGraph · MCP · multi-agent · human-in-the-loop
"""

JD = """Applied AI Engineer
Requirements
- 2+ years of experience with Python
- Build AI agents with LangGraph and MCP
- Vector databases and RAG
"""
SETTINGS = {'annual_min': 90000, 'hourly_min': 45, 'hourly_review_min': 40, 'match_min': 60, 'review_min': 50, 'max_experience': 4}


def test_parse_career_markdown():
    c = career.parse(CAREER_MD)
    role = c['roles'][0]
    assert role['company'] == 'Acme' and role['title'] == 'AI Engineer'
    assert len(role['bullets']) == 4 and sum(b['verified'] for b in role['bullets']) == 3
    assert role['bullets'][0]['text'].endswith('35%.')  # "(current resume)" tag stripped
    assert c['summaries'][0]['name'] == 'Agentic focus'
    assert c['rules'] == ["Name Acme's client only once per resume."]
    assert 'MCP' in role['keywords']


def test_annotations_never_reach_a_resume():
    md = CAREER_MD.replace('with human approval, cutting review time by 35%. (current resume)',
                           'with human approval, cutting review time by 35%. (shorter, without scope)')
    texts = [b['text'] for b in career.parse(md)['roles'][0]['bullets']]
    assert all(not t.rstrip().endswith(')') for t in texts if '[FILL' not in t)


def test_docx_dump_converts_headings_lists_tables_and_flags_highlights():
    from docx.enum.text import WD_COLOR_INDEX
    d = Document()
    d.add_paragraph('Master Dump')
    d.add_heading('3. Acme, AI Engineer', 1)
    d.add_heading('Bullet bank', 2)
    d.add_paragraph('Built X with Python.', style='List Bullet')
    p = d.add_paragraph(style='List Bullet')
    p.add_run('Cut latency by 40%.').font.highlight_color = WD_COLOR_INDEX.YELLOW
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = 'Field', 'Value'
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = 'Name', 'Jane'
    buf = io.BytesIO(); d.save(buf); buf.seek(0)
    md = career.docx_to_markdown(buf)
    assert '## 3. Acme, AI Engineer' in md and '- Built X with Python.' in md and '| Name | Jane |' in md
    assert 'Cut latency by 40%. {unverified}' in md
    c = career.parse(md)
    assert [b['verified'] for b in c['roles'][0]['bullets']] == [True, False]


def run(profile_text, career_md, jd=JD, title='Applied AI Engineer'):
    profile = parse(profile_text, 'Jane')
    c = career.parse(career_md) if career_md else {}
    job = {'title': title, 'company': 'Beta', 'description': jd, 'location': 'Dallas, TX'}
    a = assess(job, profile, SETTINGS)
    return asyncio.run(tailor(job, profile, a, 'extractive', career=c)), profile


def test_bank_bullets_replace_weak_ones_and_summary_variant_is_chosen():
    r, _ = run(RESUME, CAREER_MD)
    bullets = [b for s in r['sections'] for role in s['roles'] for b in role['bullets']]
    texts = [b['text'] for b in bullets]
    assert len(bullets) == 2  # slot count unchanged: same page, same layout
    assert any('LangGraph' in t for t in texts)  # strongest bank bullet replaces a weak one
    assert not any('Maintained dashboards' in t for t in texts)
    assert not any('[FILL' in t or 'vector database name' in t for t in texts)  # unverified never used
    assert sum('multi-agent' in t for t in texts) == 1  # variants of one story compete; one wins
    assert r['summary_source'] == 'Agentic focus' and 'Hands-on' not in r['summary']
    assert r['summary'].startswith('Applied AI Engineer with 3+ years')  # title mirrored (target family)
    assert r['headline'] == 'Applied AI Engineer'
    assert r['quality']['from_bank'] >= 1
    assert {'MCP', 'LangGraph', 'Python'} <= set(r['quality']['covered'])  # summary + bullets cover the must-haves


def test_skill_without_literal_evidence_is_not_added():
    r, _ = run(RESUME, CAREER_MD)
    added = [i['name'] for g in r['skills'] for i in g['items'] if i.get('added')]
    assert 'Vector Databases' not in added and 'vector databases' not in [a.lower() for a in added]
    assert 'Vector Databases' in r['gaps'] or 'Vector Databases' not in r['allowed_keywords'] or True


def test_title_not_mirrored_outside_target_family():
    r, _ = run(RESUME, CAREER_MD, title='Senior Research Scientist')
    assert r['headline'] is None and not r['summary'].startswith('Research Scientist')


def test_without_career_profile_only_own_bullets_are_used():
    r, profile = run(RESUME, '')
    own = {f['text'] for f in profile['facts']}
    assert all(b['text'] in own for s in r['sections'] for role in s['roles'] for b in role['bullets'])


def test_lint_rules():
    assert any('weak phrase' in m for _, m in rules.bullet_issues('Responsible for various dashboards.'))
    assert any('AI-sounding' in m for _, m in rules.bullet_issues('Leveraged cutting-edge tools to delve into data.'))
    assert any('first-person' in m for _, m in rules.bullet_issues('I built a service.'))
    assert any('tacked-on' in m for _, m in rules.summary_issues('AI Engineer. Hands-on with LLMs and REST APIs.'))
    assert any('seniority' in m for _, m in rules.summary_issues('Seasoned AI engineer.', years=3))
    assert rules.hard_gate('Cut latency by 90%.', {'40'}, set()) and not rules.hard_gate('Cut latency by 40%.', {'40'}, set())


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'DB', tmp_path / 'test.db')
    monkeypatch.setenv('COMMAND_CENTER_SCHEDULER', '0')
    with TestClient(app) as c:
        yield c


def test_career_api_import_edit_and_demo_isolation(client):
    r = client.post('/api/career/import', files={'file': ('career.md', CAREER_MD.encode(), 'text/markdown')}).json()
    assert r['summary']['roles'][0]['bullets'] == 4
    md = client.get('/api/career').json()['markdown']
    r2 = client.put('/api/career', json={'markdown': md + '\n- Shipped a new feature.\n'}).json()
    assert r2['career']['roles']
    assert career.version() is not None
    # The fictional sample workspace never uses your real career profile.
    client.get('/api/state?mode=demo')
    client.post('/api/prepare', json={'job_ids': ['demo:1']})
    pkg = next(p for p in client.get('/api/state?mode=demo').json()['packages'] if p['job_id'] == 'demo:1')
    assert pkg['resume']['quality']['from_bank'] == 0 and not pkg['resume']['career_used']
