"""Tailor by editing your own Word resume in place.

Every line of the confirmed profile is mapped back to its paragraph in the base
DOCX. Only the words that tailoring changed are touched: deleted words are cut
from their runs, inserted words become new runs that copy the formatting of the
text around them. Section order, fonts, spacing, tabs and bullets stay as you
built them. With highlight=True the inserted text is marked yellow, for review.
"""
import copy
import io
import re
from difflib import SequenceMatcher

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .resume import BULLET

TOKEN = re.compile(r"\s+|\w+(?:[.'’/+#-]\w+)*[+#]*|[^\w\s]")


def tokens(text):
    return TOKEN.findall(text or '')


def norm(text):
    text = (text or '').replace(' ', ' ').replace('’', "'")
    return re.sub(r'\s+', ' ', BULLET.sub('', text)).strip()


def segments(old, new):
    """[(text, changed)] describing `new` relative to `old` at word level."""
    a, b = tokens(old), tokens(new)
    out = []
    for op, i1, i2, j1, j2 in SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        text = ''.join(b[j1:j2])
        if not text:
            continue
        changed = op != 'equal' and bool(text.strip())
        if out and out[-1][1] == changed:
            out[-1] = (out[-1][0] + text, changed)
        else:
            out.append((text, changed))
    return out


def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _run_spans(p):
    pos, spans = 0, []
    for r in p.runs:
        spans.append((r, pos, pos + len(r.text)))
        pos += len(r.text)
    return spans


def _delete(p, start, end):
    for r, s, e in _run_spans(p):
        if e <= start or s >= end:
            continue
        a, b = max(start, s) - s, min(end, e) - s
        r.text = r.text[:a] + r.text[b:]


def _insert(p, pos, text, highlight):
    spans = _run_spans(p)
    if not spans:
        run = p.add_run(text)
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return
    # Continue the formatting of the text just before the insertion point.
    target = next(((r, s, e) for r, s, e in spans if s < pos <= e), spans[0])
    run, s, _ = target
    k = pos - s
    left, right = run.text[:k], run.text[k:]
    new_r = copy.deepcopy(run._r)
    run._r.addnext(new_r)
    new = Run(new_r, p)
    new.text = text
    new.font.highlight_color = WD_COLOR_INDEX.YELLOW if highlight else None
    if right:
        tail_r = copy.deepcopy(run._r)
        new_r.addnext(tail_r)
        Run(tail_r, p).text = right
    run.text = left


def edit_paragraph(p, new_text, highlight=False):
    """Apply word-level edits so the paragraph reads `new_text`. Returns False if it can't be done safely."""
    old = ''.join(r.text for r in p.runs)
    if not p.runs or old != p.text:  # hyperlinks/fields: text outside runs; don't risk it
        return False
    lead = old[:len(old) - len(old.lstrip())]
    a, b = tokens(old), tokens(lead + new_text.strip() if lead and not new_text.startswith(lead) else new_text)
    offsets = [0]
    for t in a:
        offsets.append(offsets[-1] + len(t))
    for op, i1, i2, j1, j2 in reversed(SequenceMatcher(None, a, b, autojunk=False).get_opcodes()):
        if op == 'equal':
            continue
        start, end = offsets[i1], offsets[i2]
        if end > start:
            _delete(p, start, end)
        ins = ''.join(b[j1:j2])
        if ins:
            _insert(p, start, ins, highlight and bool(ins.strip()))
    return True


def append_to_paragraph(p, text, highlight=False):
    current = ''.join(r.text for r in p.runs)
    if not p.runs or current != p.text:
        return False
    stripped = current.rstrip()
    pos = len(stripped) - 1 if stripped.endswith('.') else len(stripped)
    _insert(p, pos, text, highlight)
    return True


def map_facts(paragraphs, profile):
    """fact id -> paragraph, matched in document order (handles repeated lines)."""
    texts = [norm(p.text) for p in paragraphs]
    mapping, cursor = {}, 0
    for f in profile.get('facts', []):
        target = norm(f['text'])
        if not target:
            continue
        idx = next((i for i in range(cursor, len(texts)) if texts[i] == target), None)
        if idx is None:
            idx = next((i for i, t in enumerate(texts) if t == target and i not in mapping.values()), None)
        if idx is not None:
            mapping[f['id']] = idx
            cursor = idx + 1
    return mapping


def build(path, profile, resume, highlight=False):
    doc = Document(path)
    paragraphs = list(all_paragraphs(doc))
    mapping = map_facts(paragraphs, profile)
    facts = {f['id']: f for f in profile.get('facts', [])}
    report = {'mapped': len(mapping), 'total': len(facts), 'edits': [], 'failed': []}

    def para(fid):
        return paragraphs[mapping[fid]] if fid in mapping else None

    # Summary
    summary_ids = [f['id'] for f in profile['facts'] if f['section'] == 'summary']
    original = ' '.join(facts[i]['text'] for i in summary_ids)
    new_summary = (resume.get('summary') or '').strip()
    if summary_ids and new_summary and norm(new_summary) != norm(original):
        first = para(summary_ids[0])
        if first is None:
            report['failed'].append('summary: paragraph not found in your Word file')
        elif len(summary_ids) == 1:
            ok = edit_paragraph(first, new_summary, highlight)
            report['edits' if ok else 'failed'].append('summary')
        elif new_summary.endswith(original):
            ok = edit_paragraph(first, new_summary[:len(new_summary) - len(original)] + facts[summary_ids[0]]['text'], highlight)
            report['edits' if ok else 'failed'].append('summary (lead sentence)')
        else:
            ok = edit_paragraph(first, new_summary, highlight)
            for fid in summary_ids[1:]:
                p = para(fid)
                if ok and p is not None:
                    p._p.getparent().remove(p._p)
            report['edits' if ok else 'failed'].append('summary')

    # Experience / project bullets (order is kept exactly as in your document)
    for section in resume.get('sections', []):
        for role in section['roles']:
            for b in role['bullets']:
                if b['text'].strip() == b['original'].strip():
                    continue
                p = para(b['id'])
                ok = p is not None and edit_paragraph(p, b['text'], highlight)
                report['edits' if ok else 'failed'].append(f'bullet {b["id"]}')

    # Skills: each existing line keeps its label and position; items are reordered by job
    # relevance and evidence-backed additions are placed with them. No line is ever added.
    skill_paras = [mapping[f['id']] for f in profile['facts'] if f['section'] == 'skills' and f['id'] in mapping]
    for group in resume.get('skills', []):
        added = [i['name'] for i in group['items'] if i.get('added')]
        if not added and not group.get('reordered'):
            continue
        p = para(group.get('fact_id')) if group.get('fact_id') else None
        if p is None:  # older packages: find the Skills line by its label
            label = norm(group['category']).lower()
            p = next((paragraphs[i] for i in skill_paras if norm(paragraphs[i].text).lower().startswith(label + ':')), None)
        if p is not None:
            head = p.text.split(':', 1)[0] + ':' if ':' in p.text else ''
            sep = ' ' if p.text[len(head):len(head) + 1] == ' ' else ''
            ok = edit_paragraph(p, head + sep + ', '.join(i['name'] for i in group['items']), highlight)
        elif skill_paras:
            ok = append_to_paragraph(paragraphs[max(skill_paras)], ', ' + ', '.join(added), highlight)
        else:
            ok = False
        report['edits' if ok else 'failed'].append(f'skills: {group["category"]}')

    # Headline under your name ("AI Engineer"): mirror the posting's title when it's in your target family.
    if resume.get('headline'):
        for p in paragraphs[:5]:
            t = norm(p.text)
            if t and len(t.split()) <= 5 and '|' not in t and '@' not in t and re.search(r'\b(engineer|developer)\b', t, re.I):
                if t.lower() != resume['headline'].lower():
                    ok = edit_paragraph(p, resume['headline'], highlight)
                    report['edits' if ok else 'failed'].append('headline')
                break

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), report


def preview(data):
    """Paragraphs and runs of a DOCX for an HTML lookalike; highlighted runs are the tailored text."""
    doc = Document(io.BytesIO(data))
    out = []
    for p in all_paragraphs(doc):
        style_bold = bool(p.style is not None and p.style.font is not None and p.style.font.bold)
        runs = []
        for r in p.runs:
            if not r.text:
                continue
            size = r.font.size.pt if r.font.size else None
            runs.append({'text': r.text, 'bold': bool(r.bold if r.bold is not None else style_bold), 'italic': bool(r.italic),
                         'added': r.font.highlight_color == WD_COLOR_INDEX.YELLOW, 'size': size})
        align = p.alignment if p.alignment is not None else (p.style.paragraph_format.alignment if p.style is not None else None)
        is_list = p._p.pPr is not None and p._p.pPr.numPr is not None or 'list' in ((p.style.name if p.style is not None else '') or '').lower()
        out.append({'runs': runs, 'list': bool(is_list), 'center': str(align) in ('CENTER (1)', '1') or getattr(align, 'value', None) == 1,
                    'empty': not p.text.strip()})
    return out


def plain_text(data):
    doc = Document(io.BytesIO(data))
    return '\n'.join(p.text for p in all_paragraphs(doc) if p.text.strip())
